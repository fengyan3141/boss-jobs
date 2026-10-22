#!/usr/bin/env python3
"""BOSS 直聘岗位采集：边扫列表边按规则筛，筛中的立刻抓 JD。用法见 README.md。"""
import argparse
import collections
import json
import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import cities  # noqa: E402
import rules as rules_mod  # noqa: E402
import store  # noqa: E402

DEFAULT_BROWSER = Path.home() / ".boss-jobs" / "browser"


def db_path(cfg_path, cfg):
    if cfg.get("db"):
        return Path(os.path.expanduser(cfg["db"]))
    data = Path(os.environ.get("BOSS_DATA_DIR", Path.home() / ".boss-jobs" / "data"))
    data.mkdir(parents=True, exist_ok=True)
    return data / (Path(cfg_path).stem + ".sqlite")


def split(value):
    return [v.strip() for v in value.split(",") if v.strip()] if value else None


# ---- collect / details ----

def cmd_collect(args, cfg, rules, conn):
    from collector import Stop, nap

    city_list = cities.resolve(split(args.cities) or cfg.get("cities"))
    keywords = split(args.keywords) or cfg.get("keywords") or []
    if not city_list or not keywords:
        sys.exit("配置里 cities 和 keywords 都不能为空")
    print(f"方向：{rules.name}｜{len(city_list)} 城 × {len(keywords)} 词｜"
          f"{'只扫列表' if args.no_detail else '筛中即抓详情'}｜操作上限 {args.max_actions or '不限'}")

    br = open_browser(args)
    stats = collections.Counter()
    try:
        br.open_home(keywords[0], city_list[0][1])
        confirm_login(args)
        for cname, ccode in city_list:
            for kw in keywords:
                where = f"{cname}/{kw}"

                def on_batch(jobs, batch, kw=kw, where=where):
                    new, todo = 0, []
                    for raw in jobs:
                        job = store.job_from_raw(raw)
                        if not job["job_id"]:
                            continue
                        is_new = store.get(conn, job["job_id"]) is None
                        row = store.upsert_listing(conn, job, kw, rules)
                        if is_new:
                            new += 1
                            stats["新岗位"] += 1
                            stats["要详情" if row["want_detail"] else "不要"] += 1
                        if not args.no_detail and store.needs_detail(row):
                            todo.append(row)
                    got = open_details(br, conn, rules, todo, stats, args.max_details)
                    print(f"[{where} 第{batch}批] {len(jobs)} 条，新 {new}，开详情 {len(todo)}"
                          f"（{got}）｜累计 新岗位 {stats['新岗位']}，详情 {stats['详情']}，操作 {br.actions}")
                    return new

                br.search(ccode, kw, on_batch, where)
                nap(4, 8)
    except Stop as e:
        finish(stats, e)
        return
    except KeyboardInterrupt:
        finish(stats, None, "手动停止")
        return
    finish(stats, None, "全部城市和关键词跑完")


def cmd_details(args, cfg, rules, conn):
    """给库里「要详情但还没抓到」的岗位补抓（比如上次中途停了）。"""
    from collector import Stop

    todo = store.pending_details(conn, args.max_details)
    print(f"待补抓详情 {len(todo)} 条（按首次看到时间从新到旧）")
    if not todo:
        return
    br = open_browser(args)
    stats = collections.Counter()
    try:
        br.open_home("产品经理", cities.CITY_CODES["北京"])
        confirm_login(args)
        for i in range(0, len(todo), 10):
            chunk = todo[i:i + 10]
            got = open_details(br, conn, rules, chunk, stats, None)
            print(f"[{i + len(chunk)}/{len(todo)}] {got}｜操作 {br.actions}")
    except Stop as e:
        finish(stats, e)
        return
    except KeyboardInterrupt:
        finish(stats, None, "手动停止")
        return
    finish(stats, None, "补抓完成")


def open_browser(args):
    from collector import Browser
    profile = Path(os.path.expanduser(args.browser_profile))
    profile.mkdir(parents=True, exist_ok=True)
    return Browser(str(profile), args.port, args.min_interval_ms, args.max_actions)


def confirm_login(args):
    if args.yes:
        return
    print("\n浏览器已打开 BOSS。确认已经登录、能看到岗位列表后，按回车开始（Ctrl+C 退出）")
    input()


def open_details(br, conn, rules, rows, stats, max_details):
    from collector import Stop
    result = collections.Counter()
    for row in rows:
        if max_details and stats["详情"] >= max_details:
            raise Stop(f"已抓满 {max_details} 条详情上限，主动停止")
        status, jd, src = br.detail(row["job_id"])
        store.save_detail(conn, row["job_id"], status, jd, src, rules.hits(jd) if jd else "")
        stats["详情"] += 1
        stats["详情-" + status] += 1
        result[status] += 1
        if status == "ok":
            print(f"   ✓ {row['title']} / {row['company']}（{len(jd)} 字）")
    names = {"ok": "成功", "closed": "已下架", "failed": "失败"}
    return " ".join(f"{names[k]} {v}" for k, v in result.items()) or "无"


def finish(stats, stop, reason=""):
    print("\n" + "=" * 50)
    if stop is not None:
        reason = str(stop)
    print(f"结束：{reason}")
    print(f"新岗位 {stats['新岗位']}（要详情 {stats['要详情']}，不要 {stats['不要']}）｜"
          f"详情 {stats['详情']}：成功 {stats['详情-ok']}，已下架 {stats['详情-closed']}，失败 {stats['详情-failed']}")
    if stop is not None and stop.risk:
        print("⚠ 这是平台风控或登录问题。先在浏览器里手动看一眼、处理掉，隔一段时间再跑；不要马上重跑。")
    print("数据都已存进本地库，再跑会自动跳过已处理的岗位。浏览器保持打开，不用每次重新登录。")
    print("注意：这只是采集时能看到的样本，不等于市场全量。")


# ---- 离线命令 ----

def cmd_import(args, cfg, rules, conn):
    """导入旧的 jsonl（BOSS 原始列表字段，可带 jd_full），方便沿用以前抓的数据。"""
    n = jd = 0
    for path in args.files:
        with open(path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                rec = json.loads(line)
                raw = rec.get("raw", rec)
                job = store.job_from_raw(raw)
                if not job["job_id"]:
                    continue
                kw = (rec.get("_meta") or {}).get("kw", "")
                store.upsert_listing(conn, job, kw, rules)
                n += 1
                jd_text = store.normalize_text(rec.get("jd_full"))
                if jd_text:
                    store.save_detail(conn, job["job_id"], "ok", jd_text, "import", rules.hits(jd_text))
                    jd += 1
    print(f"导入 {n} 行（其中带 JD {jd}），库里现在 {count(conn)} 个岗位")


def cmd_reclassify(args, cfg, rules, conn):
    print(f"已按当前规则重新分桶 {store.reclassify(conn, rules)} 个岗位")
    cmd_stats(args, cfg, rules, conn)


def cmd_stats(args, cfg, rules, conn):
    print(f"库：{args.db}｜共 {count(conn)} 个岗位")
    rows = conn.execute(
        """select bucket, want_detail, count(*) n,
                  sum(detail_status='ok') ok, sum(detail_status='closed') closed,
                  sum(detail_status='failed') failed
           from jobs group by bucket order by want_detail desc, n desc""").fetchall()
    for r in rows:
        extra = f"  JD {r['ok']}｜下架 {r['closed']}｜失败 {r['failed']}" if r["want_detail"] else ""
        print(f"  {'●' if r['want_detail'] else '○'} {r['bucket']}：{r['n']}{extra}")
    print("  ● 开详情  ○ 不开")


def count(conn):
    return conn.execute("select count(*) from jobs").fetchone()[0]


def export_json(out, cfg, conn):
    """给 viewer/jd-viewer.html 导入用：只要已抓到 JD 的岗位。"""
    rows = conn.execute("select * from jobs where detail_status='ok' order by city, title").fetchall()
    jobs = [{
        "job_id": r["job_id"], "title": r["title"], "company": r["company"], "city": r["city"],
        "district": r["district"], "industry": r["industry"], "salary_desc": r["salary"],
        "experience": r["experience"], "degree": r["degree"],
        "skills": json.loads(r["skills"] or "[]"), "bucket": r["bucket"],
        "jd_full": store.normalize_text(r["jd"]),
    } for r in rows]
    collected = max((r["jd_at"][:10] for r in rows), default="")
    out.write_text(json.dumps({"name": cfg.get("name", ""), "collected_at": collected, "jobs": jobs},
                              ensure_ascii=False), encoding="utf-8")
    print(f"已导出 {len(jobs)} 个带 JD 的岗位 → {out}（用 viewer/jd-viewer.html 打开导入）")


def cmd_export(args, cfg, rules, conn):
    if args.out and args.out.lower().endswith(".json"):
        return export_json(Path(args.out), cfg, conn)
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font

    def clean(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

    rows = [dict(r) for r in conn.execute("select * from jobs order by city, title")]
    order = [b["name"] for b in rules.buckets]
    order += sorted({r["bucket"] for r in rows} - set(order))
    notes = {b["name"]: b["note"] for b in rules.buckets}
    city_names = sorted({r["city"] for r in rows})
    status_names = {"ok": "已抓到", "closed": "已下架", "failed": "失败", "": ""}

    wb = Workbook()
    ws = wb.active
    ws.title = "总览"
    ws.append(["分类", "开详情", "条数", "已抓到JD", "已下架", "说明"] + city_names)
    by = collections.defaultdict(list)
    for r in rows:
        by[r["bucket"]].append(r)
    for b in order:
        rs = by.get(b, [])
        ws.append([b, "是" if any(r["want_detail"] for r in rs) else "", len(rs),
                   sum(r["detail_status"] == "ok" for r in rs),
                   sum(r["detail_status"] == "closed" for r in rs), notes.get(b, "")]
                  + [sum(r["city"] == c for r in rs) for c in city_names])
    ws.append(["合计", "", len(rows)])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["F"].width = 50

    cols = ["分类", "城市", "职位", "公司", "行业", "薪资", "经验", "学历", "技能标签", "命中词",
            "搜索词", "详情", "JD命中词", "JD", "链接", "首次看到"]
    widths = [20, 6, 36, 20, 12, 12, 9, 6, 36, 24, 18, 8, 24, 80, 20, 18]
    used = set()
    for b in order:
        if not by.get(b):
            continue
        name = "".join(ch for ch in b if ch not in '[]:*?/\\')[:31]
        while name in used:
            name = name[:29] + str(len(used))
        used.add(name)
        s = wb.create_sheet(name)
        s.append(cols)
        for r in by[b]:
            s.append([clean(v) for v in (
                r["bucket"], r["city"], r["title"], r["company"], r["industry"], r["salary"],
                r["experience"], r["degree"], "、".join(json.loads(r["skills"] or "[]")), r["hits"],
                r["keyword"], status_names.get(r["detail_status"], r["detail_status"]), r["jd_hits"],
                r["jd"], f"https://www.zhipin.com/job_detail/{r['job_id']}.html", r["first_seen"])])
        for c in s[1]:
            c.font = Font(bold=True)
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions
        for i, w in enumerate(widths):
            s.column_dimensions[chr(65 + i)].width = w

    out = Path(args.out) if args.out else args.db.with_suffix(".xlsx")
    wb.save(out)
    print(f"已导出 {len(rows)} 个岗位 → {out}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def add(name, fn, help_):
        p = sub.add_parser(name, help=help_)
        p.add_argument("config", help="方向配置文件，如 profiles/aigc-pm.yaml")
        p.set_defaults(fn=fn)
        return p

    def browser_opts(p):
        p.add_argument("--max-details", type=int, help="本次最多开多少条详情")
        p.add_argument("--max-actions", type=int, default=600,
                       help="本次最多多少次页面操作（导航+滚动），默认 600；0 表示不限")
        p.add_argument("--min-interval-ms", type=int, default=800, help="两次操作最小间隔，默认 800ms")
        p.add_argument("--browser-profile", default=os.environ.get("BOSS_BROWSER_PROFILE", str(DEFAULT_BROWSER)),
                       help="Chrome 资料目录（保存登录状态），默认 ~/.boss-jobs/browser")
        p.add_argument("--port", type=int, default=int(os.environ.get("BOSS_PORT", "9333")))
        p.add_argument("--yes", action="store_true", help="不等回车确认登录（已登录、无人值守时用）")

    p = add("collect", cmd_collect, "扫列表，筛中的立刻抓详情")
    p.add_argument("--cities", help="覆盖配置，逗号分隔")
    p.add_argument("--keywords", help="覆盖配置，逗号分隔")
    p.add_argument("--no-detail", action="store_true", help="只扫列表、分桶，不开详情")
    browser_opts(p)
    browser_opts(add("details", cmd_details, "给库里待抓详情的岗位补抓"))
    p = add("import", cmd_import, "导入旧 jsonl")
    p.add_argument("files", nargs="+")
    add("reclassify", cmd_reclassify, "改了规则后，对库里所有岗位重新分桶")
    add("stats", cmd_stats, "看库里各分类数量和详情进度")
    p = add("export", cmd_export, "导出 xlsx；--out 以 .json 结尾时导出给 JD 阅读页导入的 JSON")
    p.add_argument("--out")

    args = ap.parse_args()
    cfg, rules = rules_mod.load(args.config)
    args.db = db_path(args.config, cfg)
    conn = store.connect(args.db)
    args.fn(args, cfg, rules, conn)


if __name__ == "__main__":
    main()
